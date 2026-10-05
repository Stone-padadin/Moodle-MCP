"""Read-only Moodle MCP with bounded local resource downloads."""

from __future__ import annotations

import io
import json
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

import fitz
from docx import Document
from mcp.server.fastmcp import FastMCP
from mcp.types import ToolAnnotations
from pptx import Presentation


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_BATCH_BYTES = 500 * 1024 * 1024
MAX_BATCH_FILES = 100
READ_ONLY = ToolAnnotations(
    readOnlyHint=True, destructiveHint=False, idempotentHint=True, openWorldHint=True
)
LOCAL_DOWNLOAD = ToolAnnotations(
    readOnlyHint=False, destructiveHint=False, idempotentHint=True, openWorldHint=True
)
mcp = FastMCP(
    "moodle",
    instructions=(
        "Read-only access to the user's configured Moodle site. Use these tools to read "
        "courses, course content, files, assignments and calendar deadlines. Resource "
        "downloads only create local files and never change Moodle. Never submit, upload, "
        "edit or delete Moodle content. Authentication credentials are stored locally. "
        "Moodle sites may disable Web Services or individual functions."
    ),
)


class _Text(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def _plain(value: str | None) -> str:
    parser = _Text()
    parser.feed(value or "")
    return " ".join(" ".join(parser.parts).split())


def _config_path() -> Path:
    override = os.environ.get("MOODLE_MCP_CONFIG")
    if override:
        return Path(override).expanduser()
    if os.name == "nt":
        return Path(os.environ.get("APPDATA", Path.home() / "AppData/Roaming")) / "moodle-mcp" / "config.json"
    return Path.home() / ".config" / "moodle-mcp" / "config.json"


def _config() -> dict[str, str]:
    path = _config_path()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise ValueError(
            f"Moodle is not configured. Create the local config file at {path}; see the README."
        ) from None
    except (OSError, json.JSONDecodeError):
        raise ValueError("The local Moodle config file is unreadable or invalid JSON.") from None
    base_url = str(data.get("base_url", "")).strip().rstrip("/")
    token = str(data.get("token", "")).strip()
    parsed = urllib.parse.urlsplit(base_url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("base_url must be an HTTPS Moodle site URL without embedded credentials.")
    if parsed.query or parsed.fragment:
        raise ValueError("base_url must not contain a query string or fragment.")
    if not token:
        raise ValueError("The local Moodle config has no Web Services token.")
    return {"base_url": base_url, "host": parsed.netloc.lower(), "token": token}


def _request(function: str, params: dict | None = None):
    config = _config()
    body = {
        "wstoken": config["token"],
        "wsfunction": function,
        "moodlewsrestformat": "json",
        **(params or {}),
    }
    request = urllib.request.Request(
        f"{config['base_url']}/webservice/rest/server.php",
        data=urllib.parse.urlencode(body, doseq=True).encode(),
        headers={"Content-Type": "application/x-www-form-urlencoded", "User-Agent": "moodle-mcp/1.0"},
    )
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            result = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        raise ValueError(f"Moodle API returned HTTP {exc.code}. Check the site URL and Web Services token.") from None
    except Exception:
        raise ValueError("Moodle API request failed. Check the local configuration, network, and site availability.") from None
    if isinstance(result, dict) and (result.get("exception") or result.get("errorcode")):
        message = _plain(str(result.get("message") or result.get("errorcode")))[:500]
        raise ValueError(f"Moodle API error: {message}")
    return result


def _courses() -> list[dict]:
    info = _request("core_webservice_get_site_info")
    return _request("core_enrol_get_users_courses", {"userid": info["userid"]})


def _course(course_id: int) -> dict:
    course = next((item for item in _courses() if item.get("id") == course_id), None)
    if not course:
        raise ValueError("Course is not visible in this Moodle account.")
    return course


def _public_url(url: str | None) -> str | None:
    if not url:
        return None
    parsed = urllib.parse.urlsplit(url)
    config = _config()
    if parsed.scheme != "https" or (parsed.netloc.lower() != config["host"]):
        return None
    query = [(key, value) for key, value in urllib.parse.parse_qsl(parsed.query) if key.lower() not in {"token", "wstoken"}]
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(query), ""))


class _SameHostRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        parsed = urllib.parse.urlsplit(newurl)
        if parsed.scheme != "https" or parsed.netloc.lower() != _config()["host"]:
            raise ValueError("Moodle file redirected outside the configured HTTPS site.")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def _download(file_url: str, max_bytes: int = MAX_FILE_BYTES, expected_name: str = "") -> bytes:
    config = _config()
    parsed = urllib.parse.urlsplit(file_url)
    if parsed.scheme != "https" or parsed.netloc.lower() != config["host"]:
        raise ValueError("Resource is not hosted on the configured Moodle site.")
    query = [(key, value) for key, value in urllib.parse.parse_qsl(parsed.query) if key.lower() not in {"token", "wstoken"}]
    query.append(("token", config["token"]))
    url = urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(query), ""))
    try:
        with urllib.request.build_opener(_SameHostRedirect()).open(url, timeout=90) as response:
            data = response.read(max_bytes + 1)
    except Exception:
        raise ValueError("Could not read this Moodle file. Check access and renew the token if necessary.") from None
    if len(data) > max_bytes:
        raise ValueError(f"This file exceeds the {max_bytes // (1024 * 1024)} MB download limit.")
    if Path(expected_name).suffix.lower() not in {".html", ".htm"} and data.lstrip()[:32].lower().startswith((b"<!doctype html", b"<html")):
        raise ValueError("Moodle returned a web page instead of the requested file.")
    return data


def _extension_filter(extensions: list[str] | None) -> set[str]:
    result = set()
    for value in extensions or []:
        ext = "." + value.strip().lower().lstrip(".")
        if not re.fullmatch(r"\.[a-z0-9]{1,12}", ext):
            raise ValueError(f"Invalid file extension: {value!r}")
        result.add(ext)
    return result


def _resources(course_id: int) -> list[dict]:
    sections = _request("core_course_get_contents", {"courseid": course_id})
    return [
        {
            "course_id": course_id,
            "section": section.get("name") or "Other",
            "module_id": module.get("id"),
            "module_name": module.get("name"),
            "filename": file.get("filename"),
            "mimetype": file.get("mimetype"),
            "size_bytes": file.get("filesize") or 0,
            "fileurl": file.get("fileurl"),
            "url": _public_url(module.get("url")),
        }
        for section in sections
        for module in section.get("modules", [])
        for file in module.get("contents", [])
        if file.get("type") == "file" and file.get("filename") and file.get("fileurl")
    ]


def _safe_component(value: str, max_length: int) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).replace("..", "_")
    name = name.strip(" .")[:max_length].rstrip(" .")
    reserved = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
    return "item" if not name or name.split(".", 1)[0].upper() in reserved else name


def _save_without_overwrite(root: Path, target: Path, data: bytes) -> tuple[str, Path]:
    root.mkdir(parents=True, exist_ok=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.parent.resolve().is_relative_to(root.resolve()):
        raise ValueError("Download folder resolves outside the configured Moodle download area.")
    candidate = target
    version = 2
    while candidate.exists() or candidate.is_symlink():
        if candidate.is_file() and not candidate.is_symlink() and candidate.read_bytes() == data:
            return "unchanged", candidate
        candidate = target.with_name(f"{target.stem}__{version}{target.suffix}")
        version += 1
    created = False
    try:
        with candidate.open("xb") as output:
            created = True
            output.write(data)
    except Exception:
        if created:
            candidate.unlink(missing_ok=True)
        raise
    return "downloaded", candidate


def _download_root() -> Path:
    override = os.environ.get("MOODLE_MCP_DOWNLOAD_DIR")
    return Path(override).expanduser() if override else Path.home() / "Downloads" / "Moodle"


@mcp.tool(description="List courses visible to the configured Moodle account.", annotations=READ_ONLY)
def list_courses() -> dict:
    base = _config()["base_url"]
    return {"courses": [{
        "id": course.get("id"), "code": course.get("shortname"), "name": course.get("fullname"),
        "visible": course.get("visible"), "url": f"{base}/course/view.php?id={course['id']}",
    } for course in _courses()]}


@mcp.tool(description="Read a course's sections, activities and available files.", annotations=READ_ONLY)
def get_course_outline(course_id: int) -> dict:
    course = _course(course_id)
    sections = _request("core_course_get_contents", {"courseid": course_id})
    base = _config()["base_url"]
    return {"course": {"id": course_id, "name": course.get("fullname"), "url": f"{base}/course/view.php?id={course_id}"},
        "sections": [{"name": section.get("name"), "summary": _plain(section.get("summary"))[:3000],
            "modules": [{"id": module.get("id"), "name": module.get("name"), "type": module.get("modname"),
                "description": _plain(module.get("description"))[:3000], "url": _public_url(module.get("url")),
                "files": [{"name": item.get("filename"), "type": item.get("mimetype"), "size": item.get("filesize")}
                    for item in module.get("contents", [])]}
                for module in section.get("modules", [])]} for section in sections]}


@mcp.tool(description="List downloadable files in an enrolled course; use the preview before downloading.", annotations=READ_ONLY)
def list_course_resources(course_id: int, extensions: list[str] | None = None, section_query: str = "") -> dict:
    _course(course_id)
    wanted = _extension_filter(extensions)
    resources = [{key: value for key, value in resource.items() if key != "fileurl"}
        for resource in _resources(course_id)
        if (not wanted or Path(resource["filename"]).suffix.lower() in wanted)
        and (not section_query or section_query.casefold() in resource["section"].casefold())]
    return {"course_id": course_id, "total_files": len(resources), "resources": resources}


def batch_download_resources(course_ids: list[int], module_ids: list[int] | None = None,
        extensions: list[str] | None = None, section_query: str = "", max_files: int = 50,
        dry_run: bool = False) -> dict:
    if not 1 <= len(course_ids) <= 10 or len(set(course_ids)) != len(course_ids):
        raise ValueError("Provide 1-10 distinct course IDs.")
    if not 1 <= max_files <= MAX_BATCH_FILES:
        raise ValueError(f"max_files must be 1-{MAX_BATCH_FILES}.")
    wanted = _extension_filter(extensions)
    selected_modules = set(module_ids or [])
    enrolled = {course.get("id"): course for course in _courses()}
    if any(course_id not in enrolled for course_id in course_ids):
        raise ValueError("One or more courses are not visible in this Moodle account.")
    plan = [(enrolled[course_id], resource) for course_id in course_ids for resource in _resources(course_id)
        if (not selected_modules or resource["module_id"] in selected_modules)
        and (not wanted or Path(resource["filename"]).suffix.lower() in wanted)
        and (not section_query or section_query.casefold() in resource["section"].casefold())]
    if len(plan) > max_files:
        raise ValueError(f"Selection has {len(plan)} files; narrow the filters or raise max_files (up to {MAX_BATCH_FILES}).")
    estimated_bytes = sum(int(resource["size_bytes"]) for _, resource in plan)
    if estimated_bytes > MAX_BATCH_BYTES:
        raise ValueError("Selection exceeds the 500 MB batch limit; narrow the filters.")
    root = _download_root()
    if dry_run:
        return {"dry_run": True, "files": len(plan), "estimated_bytes": estimated_bytes, "destination": str(root),
            "selection": [{"course_id": course["id"], "section": resource["section"],
                "module_id": resource["module_id"], "filename": resource["filename"], "size_bytes": resource["size_bytes"]}
                for course, resource in plan]}
    results = []
    fetched_bytes = 0
    for course, resource in plan:
        item = {"course_id": course["id"], "module_id": resource["module_id"], "filename": resource["filename"]}
        try:
            remaining = MAX_BATCH_BYTES - fetched_bytes
            if remaining <= 0:
                raise ValueError("Batch byte limit reached.")
            data = _download(resource["fileurl"], min(MAX_FILE_BYTES, remaining), resource["filename"])
            fetched_bytes += len(data)
            course_dir = _safe_component(course.get("shortname") or "course", 55) + f"-{course['id']}"
            target = root / course_dir / _safe_component(resource["section"], 70) / f"{resource['module_id']}_{_safe_component(resource['filename'], 120)}"
            status, saved = _save_without_overwrite(root, target, data)
            item.update(status=status, saved_to=str(saved), size_bytes=len(data))
        except Exception as exc:
            item.update(status="failed", error=str(exc) if isinstance(exc, ValueError) else "Local save failed.")
        results.append(item)
    return {"dry_run": False, "destination": str(root), "selected": len(plan),
        "downloaded": sum(item["status"] == "downloaded" for item in results),
        "unchanged": sum(item["status"] == "unchanged" for item in results),
        "failed": sum(item["status"] == "failed" for item in results), "results": results}


@mcp.tool(description="Download selected course files to the local Moodle folder. Preview with dry_run first; existing local files are never overwritten. Moodle is not modified.", annotations=LOCAL_DOWNLOAD)
def download_course_resources(course_ids: list[int], module_ids: list[int] | None = None,
        extensions: list[str] | None = None, section_query: str = "", max_files: int = 50,
        dry_run: bool = True) -> dict:
    return batch_download_resources(course_ids, module_ids, extensions, section_query, max_files, dry_run)


@mcp.tool(description="Read text from a PDF, PowerPoint, Word or text resource identified in the course.", annotations=READ_ONLY)
def read_course_resource(course_id: int, module_id: int, filename: str = "", start_page: int = 1, page_count: int = 6) -> dict:
    _course(course_id)
    if start_page < 1 or not 1 <= page_count <= 10:
        raise ValueError("start_page must be >= 1 and page_count must be 1-10.")
    sections = _request("core_course_get_contents", {"courseid": course_id})
    module = next((item for section in sections for item in section.get("modules", []) if item.get("id") == module_id), None)
    if not module:
        raise ValueError("Module is not in this course.")
    files = [item for item in module.get("contents", []) if not filename or item.get("filename") == filename]
    if len(files) != 1:
        raise ValueError("Specify the exact filename from the course outline.")
    file = files[0]
    name = file.get("filename") or ""
    ext = Path(name).suffix.lower()
    if ext not in {".pdf", ".docx", ".pptx", ".txt"}:
        raise ValueError("Only PDF, DOCX, PPTX and TXT reading is available.")
    data = _download(file["fileurl"], expected_name=name)
    result = {"course_id": course_id, "module_id": module_id, "filename": name,
        "url": _public_url(module.get("url")), "start_page": start_page}
    try:
        if ext == ".pdf":
            document = fitz.open(stream=data, filetype="pdf")
            result.update(total_pages=len(document), pages=[{"page": i + 1, "text": document[i].get_text()[:30000]}
                for i in range(start_page - 1, min(len(document), start_page - 1 + page_count))])
        elif ext == ".pptx":
            slides = Presentation(io.BytesIO(data)).slides
            result.update(total_pages=len(slides), pages=[{"page": i + 1, "text": "\n".join(shape.text for shape in slides[i].shapes if shape.has_text_frame)[:30000]}
                for i in range(start_page - 1, min(len(slides), start_page - 1 + page_count))])
        elif ext == ".docx":
            document = Document(io.BytesIO(data))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            result.update(text=text[(start_page - 1) * 30000:start_page * 30000], unit="30,000-character chunk (not physical page)")
        else:
            text = data.decode("utf-8-sig", errors="replace")
            result.update(text=text[(start_page - 1) * 30000:start_page * 30000], unit="30,000-character chunk (not physical page)")
    except Exception:
        raise ValueError("The file could not be parsed as its advertised format.") from None
    return result


@mcp.tool(description="List assignment activities and their release, due and cutoff dates for an enrolled course. Does not submit or change anything.", annotations=READ_ONLY)
def get_course_assignments(course_id: int) -> dict:
    _course(course_id)
    data = _request("mod_assign_get_assignments", {"courseids[0]": course_id})
    base = _config()["base_url"]
    return {"course_id": course_id, "assignments": [{"id": assignment.get("id"), "name": assignment.get("name"),
        "intro": _plain(assignment.get("intro"))[:8000], "open_time": assignment.get("allowsubmissionsfromdate"),
        "due_time": assignment.get("duedate"), "cutoff_time": assignment.get("cutoffdate"),
        "url": f"{base}/mod/assign/view.php?id={assignment['cmid']}" if assignment.get("cmid") else None}
        for course in data.get("courses", []) for assignment in course.get("assignments", [])], "warnings": data.get("warnings", [])}


@mcp.tool(description="Read upcoming Moodle calendar action events and deadlines. Events only described in course documents may not appear here.", annotations=READ_ONLY)
def get_upcoming_deadlines(days: int = 60, course_id: int = 0) -> dict:
    if not 1 <= days <= 365:
        raise ValueError("days must be 1-365.")
    if course_id:
        _course(course_id)
    now = int(time.time())
    data = _request("core_calendar_get_action_events_by_timesort", {"timesortfrom": now, "timesortto": now + days * 86400})
    return {"events": [{"name": event.get("name"), "course_id": event.get("course", {}).get("id"),
        "time": event.get("timesort"), "url": _public_url(event.get("url")),
        "description": _plain(event.get("description"))[:3000]} for event in data.get("events", [])
        if not course_id or event.get("course", {}).get("id") == course_id], "warnings": data.get("warnings", [])}


if __name__ == "__main__":
    mcp.run()
